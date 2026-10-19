#!/usr/bin/env python3
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIGURES = ROOT / "figures"

WEEKS = [
    (1, "Project Blueprint Design and Planning", "Project_Blueprint"),
    (2, "Data Exploration and Preprocessing Strategy", "Data_Exploration_Preprocessing"),
    (3, "Model Development and Training Strategy", "Model_Development_Training"),
    (4, "Model Evaluation and Performance Optimization", "Model_Evaluation_Optimization"),
    (5, "Comprehensive Project Report and Lessons Learned", "Comprehensive_Project_Report"),
]
NAVY = "17324D"
ORANGE = "D85A1A"
PALE = "F3F6F8"
GRAY = "5B6770"


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar"); fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar"); fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1); run._r.append(instr); run._r.append(fld_char2)


def set_cell_margins(cell, top=120, start=160, bottom=120, end=160):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el = tcMar.find(qn(f"w:{m}"))
        if el is None: el = OxmlElement(f"w:{m}"); tcMar.append(el)
        el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa")


def format_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72)
    sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
    sec.header_distance = Inches(.35); sec.footer_distance = Inches(.35)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string("263238")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    for name, size, color in (("Title", 30, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 14, ORANGE), ("Heading 3", 12, NAVY)):
        st = styles[name]
        st.font.name = "Calibri"; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(14 if name != "Title" else 2)
        st.paragraph_format.space_after = Pt(7)
        st.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Calibri"; styles[name].font.size = Pt(10.5)
        styles[name].paragraph_format.space_after = Pt(3)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("YUVAINTERN  |  MACHINE LEARNING CASE STUDY")
    r.font.name = "Calibri"; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(GRAY)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("Historical public-data case study  •  Page ")
    rr.font.name = "Calibri"; rr.font.size = Pt(8); rr.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_field(footer)


def add_rich_runs(p, text):
    # Small inline Markdown subset: bold and inline code.
    token = re.compile(r"(\*\*.*?\*\*|`[^`]+`)")
    pos = 0
    for m in token.finditer(text):
        if m.start() > pos: p.add_run(text[pos:m.start()])
        value = m.group(0)
        if value.startswith("**"):
            rr = p.add_run(value[2:-2]); rr.bold = True
        else:
            rr = p.add_run(value[1:-1]); rr.font.name = "Consolas"; rr.font.size = Pt(9.5); rr.font.color.rgb = RGBColor.from_string(NAVY)
        pos = m.end()
    if pos < len(text): p.add_run(text[pos:])


def add_cover(doc, week, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(46); p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"WEEK {week}  /  INTERNSHIP SUBMISSION")
    r.font.name="Calibri"; r.font.size=Pt(11); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(ORANGE)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(12)
    p.add_run(title)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(25)
    r = p.add_run("VIRTUAL JUNIOR MACHINE LEARNING ANALYST INTERN")
    r.font.size=Pt(12); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(GRAY)
    # Accent rule.
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(18)
    pPr=p._p.get_or_add_pPr(); pBdr=OxmlElement("w:pBdr"); bottom=OxmlElement("w:bottom")
    bottom.set(qn("w:val"),"single"); bottom.set(qn("w:sz"),"20"); bottom.set(qn("w:space"),"1"); bottom.set(qn("w:color"),ORANGE); pBdr.append(bottom); pPr.append(pBdr)
    t=doc.add_table(rows=4, cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(1.7); t.columns[1].width=Inches(4.7)
    entries=[("PROJECT", "Pre-call bank term-deposit subscription classification"),
             ("DATASET", "UCI Bank Marketing — bank-additional-full.csv"),
             ("SCOPE", "Reproducible historical benchmark; not a live bank deployment"),
             ("PREPARED", "19 October 2026")]
    for row,(label,value) in zip(t.rows, entries):
        row.cells[0].width=Inches(1.7); row.cells[1].width=Inches(4.7)
        for cell in row.cells: set_cell_margins(cell); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(row.cells[0], NAVY); shade(row.cells[1], PALE)
        p0=row.cells[0].paragraphs[0]; p0.paragraph_format.space_after=Pt(0)
        rr=p0.add_run(label); rr.font.size=Pt(8.5); rr.font.bold=True; rr.font.color.rgb=RGBColor(255,255,255)
        p1=row.cells[1].paragraphs[0]; p1.paragraph_format.space_after=Pt(0)
        rr=p1.add_run(value); rr.font.size=Pt(10); rr.font.color.rgb=RGBColor.from_string(NAVY)
    doc.add_paragraph()
    box=doc.add_table(rows=1,cols=1); box.alignment=WD_TABLE_ALIGNMENT.CENTER
    cell=box.cell(0,0); shade(cell,"FFF4EA"); set_cell_margins(cell,top=220,bottom=220,start=240,end=240)
    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run("Project boundary  "); r.bold=True; r.font.color.rgb=RGBColor.from_string(ORANGE)
    p.add_run("Results describe a public historical dataset. They do not establish current deployment performance, causal campaign impact, or business profitability.")
    doc.add_page_break()


def add_contents(doc, markdown):
    heads=[]
    for line in markdown.splitlines():
        m=re.match(r"^(#{1,3})\s+(.*)$",line)
        if m and len(m.group(1))>=2:
            title=m.group(2).strip()
            if title.lower()!="references": heads.append((len(m.group(1)), title))
    doc.add_heading("Contents", level=1)
    for level,title in heads:
        p=doc.add_paragraph()
        p.paragraph_format.left_indent=Inches(.22*(level-2))
        p.paragraph_format.space_after=Pt(2)
        rr=p.add_run(title); rr.font.size=Pt(10 if level==2 else 9.5); rr.font.color.rgb=RGBColor.from_string(NAVY if level==2 else GRAY)
    doc.add_page_break()


def add_figures(doc, week):
    if week==2:
        images=[("target_distribution.png","Figure 1. Target distribution in the full historical dataset."),
                ("month_conversion_rate.png","Figure 2. Historical subscription rate by campaign month; descriptive, not causal.")]
    elif week==4:
        images=[("precision_recall.png","Figure 1. Precision–recall curves on the chronological holdout."),
                ("random_forest_confusion_matrix.png","Figure 2. Random-forest confusion matrix at the illustrative 0.50 threshold.")]
    else: return
    doc.add_heading("Selected Project Figures", level=1)
    for filename,caption in images:
        path=FIGURES/filename
        if path.exists():
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(1)
            p.add_run().add_picture(str(path), width=Inches(5.8))
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
            r=p.add_run(caption); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(GRAY)


def add_body(doc, markdown, week):
    lines=markdown.splitlines()
    buffer=[]
    for line in lines:
        s=line.strip()
        if not s:
            continue
        h=re.match(r"^(#{1,6})\s+(.*)$",s)
        if h:
            text=h.group(2).strip()
            if text.lower()=="references":
                add_figures(doc,week)
            level=min(len(h.group(1)),3)
            doc.add_heading(text,level=level)
            continue
        if s.startswith("---"):
            continue
        if re.match(r"^\d+\.\s+",s):
            p=doc.add_paragraph(style="List Number"); add_rich_runs(p,re.sub(r"^\d+\.\s+","",s)); continue
        if s.startswith("- "):
            p=doc.add_paragraph(style="List Bullet"); add_rich_runs(p,s[2:]); continue
        p=doc.add_paragraph()
        add_rich_runs(p,s)


def main():
    for week,title,slug in WEEKS:
        md=(REPORTS/f"week{week}_content.md").read_text(encoding="utf-8")
        doc=Document(); format_styles(doc)
        doc.core_properties.title=f"Week {week} — {title}"
        doc.core_properties.subject="YuvaIntern Virtual Junior Machine Learning Analyst Intern submission"
        doc.core_properties.author="Manus AI"
        doc.core_properties.keywords="machine learning, bank marketing, classification, YuvaIntern"
        add_cover(doc,week,title)
        add_contents(doc,md)
        add_body(doc,md,week)
        path=REPORTS/f"Week_{week}_{slug}.docx"
        doc.save(path)
        print(path)

if __name__=="__main__": main()
